import os
import tempfile
import requests
import streamlit as st
import whisper
from docx import Document

# ===================== 頁面設定 =====================
st.set_page_config(
    page_title="MP3 轉中英對照文字",
    page_icon="🎙️",
    layout="wide"
)
st.title("🎙️ MP3 語音轉文字｜中英對照Word產生器")
st.markdown("支援國語、英文混合音檔，Whisper語音辨識 + Google翻譯，輸出docx檔下載")

# ===================== 自製Google翻譯函數（無googletrans，避免httpx衝突） =====================
def google_translate(text, target_lang="en"):
    if not text.strip():
        return ""
    url = "https://translate.google.com/translate_a/single"
    params = {
        "client": "gtx",
        "sl": "auto",
        "tl": target_lang,
        "dt": "t",
        "q": text
    }
    headers = {
        "User-Agent": "Mozilla/5.0"
    }
    try:
        resp = requests.get(url, params=params, headers=headers, timeout=15)
        json_data = resp.json()
        translated_text = "".join([item[0] for item in json_data[0]])
        return translated_text
    except Exception as e:
        return f"[翻譯失敗] {str(e)}"

# ===================== 載入Whisper模型，快取 =====================
@st.cache_resource
def load_whisper_model():
    # tiny / base / small / medium，免費雲端建議 small
    return whisper.load_model("small")

model = load_whisper_model()

# ===================== 上傳音檔 =====================
uploaded_file = st.file_uploader("請上傳 MP3 音檔", type=["mp3"])

if uploaded_file is not None:
    # 建立暫存音檔
    with tempfile.NamedTemporaryFile(delete=False, suffix=".mp3") as tmp_audio:
        tmp_audio.write(uploaded_file.read())
        audio_path = tmp_audio.name

    st.info("⏳ 開始語音辨識，請稍候...")
    # Whisper辨識，自動偵測中文/英文
    result = model.transcribe(
        audio_path,
        language=None,
        task="transcribe",
        verbose=False
    )
    segments = result["segments"]
    full_raw_text = result["text"]

    st.success("✅ 語音辨識完成！")
    st.subheader("原始辨識全文")
    st.write(full_raw_text)

    # ===================== 建立Word文件 =====================
    doc = Document()
    doc.add_heading('中英對照 語音轉文字結果', level=1)
    # 建立2欄表格
    table = doc.add_table(rows=1, cols=2)
    table.style = 'Table Grid'
    header_cells = table.rows[0].cells
    header_cells[0].text = "原文（辨識結果）"
    header_cells[1].text = "英文翻譯"

    st.info("⏳ 逐段翻譯中...")
    for seg in segments:
        seg_text = seg["text"].strip()
        if not seg_text:
            continue
        trans_en = google_translate(seg_text, target_lang="en")
        row_cells = table.add_row().cells
        row_cells[0].text = seg_text
        row_cells[1].text = trans_en

    # 加入完整辨識全文區塊
    doc.add_paragraph("")
    doc.add_heading('完整辨識全文', level=2)
    doc.add_paragraph(full_raw_text)

    # 儲存docx到暫存
    tmp_docx = tempfile.NamedTemporaryFile(delete=False, suffix=".docx")
    doc.save(tmp_docx.name)
    tmp_docx.close()

    # 讀取檔案二進制，提供下載
    with open(tmp_docx.name, "rb") as f:
        doc_file_bytes = f.read()

    st.success("📄 Word文件產生完畢！點擊下方按鈕下載")
    st.download_button(
        label="📥 下載中英對照 docx",
        data=doc_file_bytes,
        file_name="audio_transcript_bilingual.docx",
        mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    )

    # 清理暫存檔
    os.unlink(audio_path)
    os.unlink(tmp_docx.name)

